import io
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "resources" / "reimbursement-blank.xlsx"


def _safe(value):
    text = str(value or "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def _copy_row(donor, sheet, source_number, target_number):
    sheet.row_dimensions[target_number].height = donor.row_dimensions[source_number].height
    for column in range(1, 7):
        source = donor.cell(source_number, column)
        target = sheet.cell(target_number, column)
        target._style = copy(source._style)
        target.protection = copy(source.protection)


def workbook_bytes(snapshot, owner):
    donor_book = load_workbook(TEMPLATE_PATH)
    donor = donor_book.active
    book = load_workbook(TEMPLATE_PATH)
    sheet = book.active
    for merged in list(sheet.merged_cells.ranges):
        if str(merged) != "A1:F1":
            sheet.unmerge_cells(str(merged))
    sheet.delete_rows(3, sheet.max_row - 2)
    _copy_row(donor, sheet, 1, 1)
    sheet["A1"] = _safe(snapshot["title"])
    _copy_row(donor, sheet, 2, 2)
    for column, value in enumerate(("负责人", "项目", "数量", "单价", "备注", "费用"), 1):
        sheet.cell(2, column, value)
    row = 3
    for index, block in enumerate(snapshot["blocks"]):
        items = []
        prices = []
        remarks = []
        for order in block["orders"]:
            items.append(f"【{order['label']}】")
            prices.append(f"【{order['label']}】")
            for line in order["lines"]:
                items.append(f"{line['name']} {line['quantity']}{line['unit']}")
                prices.append(f"{line['name']}：{line['price']}（{line['price_type']}）")
            for adjustment in order["adjustments"]:
                prices.append(f"{adjustment['label']}：{adjustment['amount']}")
            if order["expected"]:
                remarks.append(
                    f"{order['label']} 应付 {order['expected']}；实付分配 {order['paid']}"
                )
            if order["issues"]:
                remarks.append(f"{order['label']}：{'；'.join(order['issues'])}")
            if snapshot["exceptions"].get(order["id"]):
                remarks.append(f"核对说明：{snapshot['exceptions'][order['id']]}")
        _copy_row(donor, sheet, 3 if index == 0 else 6, row)
        values = [
            _safe(owner),
            block["venue"],
            "\n".join(items),
            "\n".join(prices),
            "\n".join(remarks),
            f"总计：{block['total']}元",
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, _safe(value))
        sheet.row_dimensions[row].height = min(
            400, max(70, (max(len(items), len(prices)) + 1) * 18)
        )
        row += 1
        _copy_row(donor, sheet, 4, row)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        sheet.cell(row, 1, _safe(f"负责人：{owner}"))
        sheet.cell(row, 3, f"总计：{block['total']}元")
        row += 1
        _copy_row(donor, sheet, 5, row)
        row += 1
    _copy_row(donor, sheet, 15, row)
    sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
    sheet.cell(row, 5, f"总费用：{snapshot['total']}元")
    sheet.freeze_panes = "A3"
    sheet.print_title_rows = "1:2"
    sheet.print_area = f"A1:F{row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
