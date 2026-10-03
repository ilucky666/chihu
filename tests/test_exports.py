import hashlib
import io
import zipfile

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from core.models import User
from core.services import exports, finance
from core.services.projects import create_project, create_visit
from core.services.template_workbook import TEMPLATE_PATH


@pytest.mark.django_db
def test_four_blocks_and_draft_vs_final():
    template_digest = hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
    leader = User.objects.create_user(username="export-leader")
    project = create_project(leader, "四店测评")
    visits = []
    for index, amount in enumerate((550, 122, 437, 258)):
        visit = create_visit(leader, project, f"店 {index}")
        order = finance.create_order(leader, visit, f"单 {index}", expected_amount=amount)
        finance.add_line(leader, order, f"套餐 {index}", price=amount)
        payment = finance.create_payment(leader, visit, leader, amount)
        finance.allocate(leader, payment, order, amount)
        visits.append(visit)
    claim = exports.create_claim(leader, project, [v.pk for v in visits])
    draft = exports.export_claim(leader, claim)
    assert draft.snapshot["total"] == "1367.00"
    with draft.file.open("rb") as source, zipfile.ZipFile(io.BytesIO(source.read())) as archive:
        workbook = load_workbook(io.BytesIO(archive.read("报销表.xlsx")))
        sheet = workbook.active
        assert [sheet.cell(2, col).value for col in range(1, 7)] == [
            "负责人",
            "项目",
            "数量",
            "单价",
            "备注",
            "费用",
        ]
        assert sheet["E15"].value == "总费用：1367.00元"
        assert {str(item) for item in sheet.merged_cells.ranges} == {
            "A1:F1",
            "A4:B4",
            "C4:F4",
            "A7:B7",
            "C7:F7",
            "A10:B10",
            "C10:F10",
            "A13:B13",
            "C13:F13",
            "E15:F15",
        }
        assert "索引.json" in archive.namelist()
    assert hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest() == template_digest
    with pytest.raises(ValidationError):
        exports.export_claim(leader, claim, final=True)
    snapshot, _ = exports.claim_snapshot(claim)
    claim.exception_notes = {
        o["id"]: "凭证待补，负责人已核对" for b in snapshot["blocks"] for o in b["orders"]
    }
    claim.save()
    final = exports.export_claim(leader, claim, final=True)
    assert final.status == "final"
    assert not exports.finance_is_stale(final)
    first_order = visits[0].orders.get()
    first_order.label = "修改后的订单"
    first_order.save()
    assert exports.finance_is_stale(final)
    with pytest.raises(ValidationError):
        exports.export_claim(leader, claim, final=True)


def test_formula_like_text_is_escaped():
    assert exports.safe_text('=HYPERLINK("bad")').startswith("'=")
    assert exports.safe_text("+SUM(1,2)").startswith("'+")
